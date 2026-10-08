"""Resume ingestion: file → clean lines tagged with section + project block.

Design choices
--------------
* **Extractor tournament.** Resume PDFs come from Word, LaTeX, Canva, online
  builders… No single extractor wins on all of them (two-column layouts break
  pdfplumber, letter-spaced headings break pypdf). We run both and keep the
  output with the best readability score.
* **Hyperlink harvesting.** Many resumes show a GitHub *icon* whose URL lives
  only in the PDF link annotations. We read those too.
* **Layout-agnostic sections.** Headings are recognised by vocabulary
  ("Work History", "Professional Experience", "Internships" → experience), not
  by position or font.
"""
from __future__ import annotations

import hashlib
import io
import logging
import re
import unicodedata
from pathlib import Path

from .models import Block, Line, ParsedResume

log = logging.getLogger(__name__)
for _noisy in ("pypdf", "pdfminer", "pdfplumber"):   # malformed-PDF chatter is expected, not actionable
    logging.getLogger(_noisy).setLevel(logging.ERROR)

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class ResumeReadError(Exception):
    """Raised when a file cannot be turned into usable text."""


# ── raw extraction ────────────────────────────────────────────────────────────

def _pdf_links(data: bytes) -> list[str]:
    from pypdf import PdfReader
    links: list[str] = []
    try:
        reader = PdfReader(io.BytesIO(data))
        for page in reader.pages:
            for annot in page.get("/Annots") or []:
                obj = annot.get_object()
                action = obj.get("/A")
                if action and action.get("/URI"):
                    links.append(str(action["/URI"]).strip())
    except Exception as exc:  # annotations are a bonus, never fatal
        log.debug("link extraction failed: %s", exc)
    return links


def _pdf_candidates(data: bytes) -> dict[str, str]:
    out: dict[str, str] = {}
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        out["pypdf"] = "\n".join((p.extract_text() or "") for p in reader.pages)
    except Exception as exc:
        log.debug("pypdf failed: %s", exc)
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            out["pdfplumber"] = "\n".join((p.extract_text() or "") for p in pdf.pages)
    except Exception as exc:
        log.debug("pdfplumber failed: %s", exc)
    return out


def _docx_text(data: bytes) -> tuple[str, list[str]]:
    import docx  # python-docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    links = [r.target_ref for r in doc.part.rels.values() if "hyperlink" in r.reltype]
    return "\n".join(parts), links


# ── normalisation ─────────────────────────────────────────────────────────────

_LIGATURES = {"ﬁ": "fi", "ﬂ": "fl", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl"}
_BULLETS = re.compile(r"^[\s•●▪◦○■□➢➤►\-–—*·•]+")


def _despace(line: str) -> str:
    """'S K I L L S  P y t h o n' → 'SKILLS Python' (letter-spaced PDFs)."""
    tokens = line.split(" ")
    singles = sum(1 for t in tokens if len(t) == 1)
    if len(tokens) >= 6 and singles / len(tokens) > 0.6:
        words = re.split(r"\s{2,}", line)
        return " ".join(w.replace(" ", "") for w in words)
    return line


def _reflow(lines: list[str]) -> list[str]:
    """Join word-per-line runs (justified text in some PDFs) back into sentences."""
    out: list[str] = []
    run: list[str] = []

    def flush() -> None:
        if len(run) >= 4 and out:
            out[-1] = out[-1] + " " + " ".join(run)
        else:
            out.extend(run)
        run.clear()

    for line in lines:
        if len(line.split()) <= 2 and _heading(line) is None:
            run.append(line)
        else:
            flush()
            out.append(line)
    flush()
    return out


def normalise(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    for k, v in _LIGATURES.items():
        text = text.replace(k, v)
    text = text.replace("(cid:127)", "•").replace(" ", " ")
    lines = []
    for raw in text.splitlines():
        line = _despace(raw.rstrip())
        line = _BULLETS.sub("", line).strip()
        line = re.sub(r"[ \t]{2,}", " ", line)
        if line:
            lines.append(line)
    return "\n".join(_reflow(lines))


_WORD = re.compile(r"[A-Za-z][A-Za-z+#.\-]{1,}")


def readability(text: str) -> float:
    """0..1 heuristic: share of tokens that look like real words, penalising
    letter-spacing, glued words and word-per-line output."""
    tokens = text.split()
    lines = [l for l in text.splitlines() if l.strip()]
    if len(tokens) < 12 or not lines:
        return 0.0
    wordy = sum(1 for t in tokens if _WORD.fullmatch(t.strip(",.;:()[]|")) and 2 <= len(t) <= 18)
    singles = sum(1 for t in tokens if len(t) == 1)
    glued = sum(1 for t in tokens if len(t) > 25)
    short_lines = sum(1 for l in lines if len(l.split()) <= 1) / len(lines)
    score = wordy / len(tokens) - 0.5 * singles / len(tokens) - glued / len(tokens) - 0.4 * short_lines
    return round(max(0.0, min(1.0, score)), 3)


# ── sections & blocks ─────────────────────────────────────────────────────────

_HEADINGS: list[tuple[str, re.Pattern[str]]] = [
    ("experience", re.compile(r"^(work\s+|professional\s+|relevant\s+|industry\s+)?(experience|employment|work\s+history|internships?|career\s+history)\b", re.I)),
    ("projects", re.compile(r"^(key\s+|academic\s+|personal\s+|selected\s+|technical\s+|major\s+|notable\s+)?projects?\b", re.I)),
    ("skills", re.compile(r"^(technical\s+|core\s+|key\s+)?(skills|skill\s+set|technologies|tech\s+stack|tools(\s*&\s*technologies)?|competencies|expertise)\b", re.I)),
    ("education", re.compile(r"^(education(al)?(\s+qualifications?)?|academics?|academic\s+background|relevant\s+coursework|coursework)\b", re.I)),
    ("certifications", re.compile(r"^(certifications?|certificates|licenses|courses|trainings?|achievements|awards|honou?rs|publications|extra[\s-]?curricular|co[\s-]curricular|activities|leadership|positions?\s+of\s+responsibility|volunteer\w*|languages\s+known|hobbies|interests|area\s+of\s+interest)\b", re.I)),
    ("summary", re.compile(r"^(summary|professional\s+summary|profile|about(\s+me)?|objective|career\s+(objective|aspiration)|resume\s+summary)\b", re.I)),
]


def _unkern(text: str) -> str:
    """'PROFESSIONAL SUMMAR Y' → 'PROFESSIONAL SUMMARY' (kerning splits in some PDF exports)."""
    out: list[str] = []
    for tok in text.split():
        if out and tok.isalpha() and out[-1].isalpha() and tok.isupper() and out[-1].isupper() \
                and (len(tok) <= 2 or len(out[-1]) <= 2):
            out[-1] += tok
        else:
            out.append(tok)
    return " ".join(out)


def _heading(line: str) -> str | None:
    clean = line.strip().strip(":").strip()
    if clean.isupper():
        clean = _unkern(clean)
    if len(clean) > 45 or len(clean.split()) > 6:
        return None
    for name, rx in _HEADINGS:
        if rx.match(clean):
            # 'Python Projects | Django' is a project title, not a heading
            if "|" in clean or re.search(r"\d{4}", clean):
                return None
            return name
    return None


_DATE = re.compile(r"((jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*['’]?\d{2,4}|\b\d{1,2}/\d{4}\b|\b20\d{2}\b)\s*[-–—to]+\s*(present|current|now|(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s*['’]?\d{2,4}|\d{1,2}/\d{4}|20\d{2})", re.I)


def _is_block_title(line: str, prev_was_bullet: bool) -> bool:
    words = line.split()
    if not words or len(words) > 16:
        return False
    if line.endswith(".") and not _DATE.search(line):
        return False
    if "|" in line or _DATE.search(line) or re.search(r"\b(github|live\s+link|link|demo)\b\s*$", line, re.I):
        return True
    titled = sum(1 for w in words if w[:1].isupper()) / len(words)
    return prev_was_bullet and len(words) <= 10 and titled >= 0.6


def segment(text: str) -> tuple[list[Line], list[Block], list[str]]:
    raw_lines = text.splitlines()
    lines: list[Line] = []
    blocks: list[Block] = []
    section = "summary" if raw_lines else "other"
    seen_heading = False
    found: list[str] = []
    current: Block | None = None
    prev_bullet = True
    for i, txt in enumerate(raw_lines):
        head = _heading(txt)
        if head:
            section = head
            seen_heading = True
            found.append(head)
            current = None
            prev_bullet = True
            lines.append(Line(no=i, text=txt, section=section))
            continue
        sec = section if (seen_heading or i < 12) else "other"
        if i < 6 and not seen_heading:
            sec = "summary"  # header area: name, contacts, tagline
        block_id = None
        if sec in ("experience", "projects"):
            if current is None or _is_block_title(txt, prev_bullet):
                current = Block(id=len(blocks), section=sec, title=txt[:120], lines=[], word_count=0)
                blocks.append(current)
            current.lines.append(i)
            current.word_count += len(txt.split())
            block_id = current.id
        lines.append(Line(no=i, text=txt, section=sec, block=block_id))
        prev_bullet = len(txt.split()) > 10
    text_of = {ln.no: ln.text for ln in lines}
    sec_of = {ln.no: ln.section for ln in lines}
    for b in blocks:  # readable titles: 'Aug 2024 – Present | Pune' + 'ML Intern, Acme' → combined
        first = text_of[b.lines[0]]
        core = re.sub(r"[|,·–—\-]", " ", _DATE.sub("", first)).split()
        if len(core) >= 3:
            continue
        prev = text_of.get(b.lines[0] - 1, "")
        nxt = text_of.get(b.lines[1], "") if len(b.lines) > 1 else ""
        if prev and sec_of.get(b.lines[0] - 1) == b.section and _heading(prev) is None and len(prev.split()) <= 12 \
                and not prev.endswith("."):
            b.title = (prev + " · " + first)[:120]
        elif nxt and len(nxt.split()) <= 14:
            b.title = (nxt + " · " + first)[:120]
    if not seen_heading:  # unstructured resume: treat everything as one 'other' block
        for ln in lines:
            ln.section = "other"
    return lines, blocks, sorted(set(found))


# ── contact fields ────────────────────────────────────────────────────────────

EMAIL = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
GITHUB = re.compile(r"github\.com/([A-Za-z0-9](?:[A-Za-z0-9-]{0,38}))", re.I)
_GITHUB_RESERVED = {"features", "about", "orgs", "topics", "settings", "login", "join", "sponsors", "marketplace"}
_NOT_NAME = re.compile(r"@|\d{3}|http|www\.|linkedin|github|resume|curriculum|vitae|engineer|developer|intern|"
                       r"summary|profile|objective|student|analyst|scientist", re.I)


def find_github(text: str, links: list[str]) -> str | None:
    counts: dict[str, list] = {}
    for src in [text, *links]:
        for m in GITHUB.finditer(src):
            user = m.group(1)
            if user.lower() in _GITHUB_RESERVED:
                continue
            counts.setdefault(user.lower(), [0, user])[0] += 1
    if not counts:
        return None
    return max(counts.values(), key=lambda v: v[0])[1]


_NAME_STOP = re.compile(r"\s*(?:[|•·—,:(]|\S+@|\+?\d|https?:|www\.|\b(?:e-?mail|phone|mobile|linkedin|github|"
                        r"portfolio|leetcode|address|contact)\b)", re.I)


def find_name(lines: list[str]) -> str | None:
    """First plausible person name in the header area."""
    fallback = None
    for raw in lines[:6]:
        head = _NAME_STOP.split(raw.strip(), maxsplit=1)[0].strip()
        words: list[str] = []
        for w in head.split():
            glued = re.match(r"^([A-Z][a-z]+)([A-Z][a-z]+.*)$", w)   # 'KumarDehradun' → 'Kumar'
            if glued:
                words.append(glued.group(1))
                break
            words.append(w)
        words = words[:4]
        cand = " ".join(words)
        if not words or _NOT_NAME.search(cand) or not all(re.fullmatch(r"[A-Za-z.'\-]+", w) for w in words):
            continue
        pretty = " ".join(w if (w.isupper() and len(w) <= 2) or "." in w else w.capitalize() for w in words)
        if len(words) >= 2:
            return pretty
        fallback = fallback or pretty
    return fallback


# ── public API ────────────────────────────────────────────────────────────────

def read_resume(path: Path) -> ParsedResume:
    data = path.read_bytes()
    if not data:
        raise ResumeReadError("empty file")
    sha = hashlib.sha256(data).hexdigest()
    ext = path.suffix.lower()
    links: list[str] = []
    if ext == ".pdf":
        if not data.lstrip().startswith(b"%PDF"):
            raise ResumeReadError("not a valid PDF (missing %PDF header)")
        candidates = _pdf_candidates(data)
        links = _pdf_links(data)
    elif ext == ".docx":
        try:
            raw, links = _docx_text(data)
        except Exception as exc:
            raise ResumeReadError(f"unreadable DOCX: {exc}") from exc
        candidates = {"python-docx": raw}
    elif ext in (".txt", ".md"):
        candidates = {"text": data.decode("utf-8", errors="replace")}
    else:
        raise ResumeReadError(f"unsupported file type {ext}")

    # Tournament: every extractor's output is normalised, then judged on
    # readability + how many resume sections we can recognise in it.
    entries = []
    for name, raw in candidates.items():
        if not raw or not raw.strip():
            continue
        despaced = "\n".join(_despace(l) for l in unicodedata.normalize("NFKC", raw).splitlines())
        q = readability(despaced)
        norm = normalise(raw)
        n_sections = len(segment(norm)[2])
        entries.append((q + 0.03 * n_sections, q, name, norm))
    entries.sort(key=lambda e: e[0], reverse=True)
    if not entries or entries[0][1] < 0.15:
        raise ResumeReadError("no extractable text (scanned image or corrupt file?)")
    _, quality, extractor, text = entries[0]

    lines, blocks, found = segment(text)
    raw_lines = text.splitlines()
    # mailto: annotations beat text (icon fonts often glue glyphs onto the address)
    mailto = next((l[7:].split("?")[0].strip() for l in links
                   if l.lower().startswith("mailto:") and EMAIL.fullmatch(l[7:].split("?")[0].strip())), None)
    in_text = next((m.group(0) for m in EMAIL.finditer(text)), None)
    if in_text and mailto and in_text.lower().endswith(mailto.lower()):
        email = mailto          # text had icon glyphs glued in front of the address
    else:
        email = in_text or mailto
    return ParsedResume(
        file=path.name, sha256=sha, extractor=extractor, extraction_quality=round(quality, 3), text=text,
        lines=lines, blocks=blocks, links=sorted(set(links)), name=find_name(raw_lines),
        email=email.lower() if email else None, github_username=find_github(text, links), sections_found=found,
    )


def safe_read(path: Path) -> tuple[str, ParsedResume | None, str | None]:
    """Process-pool friendly wrapper: never raises."""
    try:
        return path.name, read_resume(path), None
    except ResumeReadError as exc:
        return path.name, None, str(exc)
    except Exception as exc:  # defensive: a broken file must not kill the batch
        return path.name, None, f"{type(exc).__name__}: {exc}"


def discover(folder: Path) -> list[Path]:
    return sorted(p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED)
